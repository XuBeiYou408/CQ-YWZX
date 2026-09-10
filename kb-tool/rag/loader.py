import os
import logging
from typing import List, Tuple, Dict
from langchain_community.document_loaders import PyMuPDFLoader, TextLoader
from langchain_core.documents import Document

from config import folder_path
from rag.splitter import pdf_qingxi, md_qingxi, docx_qingxi, txt_qingxi

logger = logging.getLogger(__name__)

def _load_docx(file_path: str) -> List[Document]:
    """使用 python-docx 提取段落与表格文本"""
    try:
        from docx import Document as DocxReader
        doc = DocxReader(file_path)
        text_parts = []
        for p in doc.paragraphs:
            if p.text.strip():
                text_parts.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_parts.append(" | ".join(row_text))
        text = "\n\n".join(text_parts)
        if not text.strip():
            return []
        return [Document(page_content=text, metadata={"source": file_path, "file_type": "docx"})]
    except Exception as e:
        logger.error(f"读取 docx 失败 {file_path}: {e}")
        return []

def _load_txt(file_path: str) -> List[Document]:
    """读取 txt 纯文本"""
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            text = f.read()
        if not text.strip():
            return []
        return [Document(page_content=text, metadata={"source": file_path, "file_type": "txt"})]
    except Exception as e:
        logger.error(f"读取 txt 失败 {file_path}: {e}")
        return []

def load_and_split_single_file(file_path: str) -> List[Document]:
    """加载并切分单个文档（支持 pdf, md, docx, txt）"""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.pdf':
        docs = PyMuPDFLoader(file_path).load()
        return pdf_qingxi(docs)
    elif ext == '.md':
        docs = TextLoader(file_path, encoding='utf-8').load()
        return md_qingxi(docs)
    elif ext == '.docx':
        docs = _load_docx(file_path)
        return docx_qingxi(docs)
    elif ext == '.txt':
        docs = _load_txt(file_path)
        return txt_qingxi(docs)
    else:
        logger.warning(f"不支持的文件格式: {ext}")
        return []

def load_all_documents() -> Tuple[List[Document], List[Document], List[Document], List[Document], Dict[str, int]]:
    """扫描素材目录，分类加载所有支持的文档格式"""
    pdf_list = []
    md_list = []
    docx_list = []
    txt_list = []
    file_count = {'pdf': 0, 'md': 0, 'docx': 0, 'txt': 0}
    logger.info(f"正在扫描处理本地文件中: {folder_path} ...")
    if os.path.exists(folder_path):
        for root, _, files in os.walk(folder_path):
            for file in files:
                full_path = os.path.join(root, file)
                ext = os.path.splitext(file)[1].lower()
                if ext == '.pdf':
                    try:
                        pdf_list.extend(PyMuPDFLoader(full_path).load())
                        file_count['pdf'] += 1
                    except Exception as e:
                        logger.error(f"处理 PDF 文件 {file} 出错：{e}")
                elif ext == '.md':
                    try:
                        md_list.extend(TextLoader(full_path, encoding='utf-8').load())
                        file_count['md'] += 1
                    except Exception as e:
                        logger.error(f"处理 MD 文件 {file} 出错：{e}")
                elif ext == '.docx':
                    try:
                        docx_list.extend(_load_docx(full_path))
                        file_count['docx'] += 1
                    except Exception as e:
                        logger.error(f"处理 Word 文件 {file} 出错：{e}")
                elif ext == '.txt':
                    try:
                        txt_list.extend(_load_txt(full_path))
                        file_count['txt'] += 1
                    except Exception as e:
                        logger.error(f"处理 TXT 文件 {file} 出错：{e}")
    logger.info(f"处理完成，共处理 {file_count['pdf']} 个 PDF，{file_count['md']} 个 Markdown，{file_count['docx']} 个 Word，{file_count['txt']} 个 TXT。")
    return pdf_list, md_list, docx_list, txt_list, file_count
