import io
import sys
import os

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docx
import pymupdf
from fastapi.testclient import TestClient
from app.main import app
from contract.parser import parse_uploaded_file, detect_contract_type
from contract.precedents import search_precedents

client = TestClient(app)

def run_tests():
    print("================================================================")
    print(">> 测试 1: 构造并测试 Word (.docx) 文件内存生成与解析")
    print("================================================================")
    doc = docx.Document()
    doc.add_heading("北京市房屋租赁合同 (示范文本)", level=1)
    doc.add_paragraph("出租人（甲方）：王大明  身份证号：11010119800101xxxx")
    doc.add_paragraph("承租人（乙方）：张小红  身份证号：31010119950202xxxx")
    doc.add_paragraph("第一条 房屋基本情况：甲方将位于北京市朝阳区某某小区5号楼302室出租给乙方居住使用。")
    doc.add_paragraph("第二条 租赁期限：租赁期限自2026年10月1日起至2051年10月1日止，共计25年。") # 大坑：超20年法定上限
    doc.add_paragraph("第三条 租金及押金：押金为人民币10,000元。无论合同以何种方式解除，甲方均有权不予退还押金。") # 大坑：无理扣押金
    
    docx_bytes = io.BytesIO()
    doc.save(docx_bytes)
    docx_data = docx_bytes.getvalue()
    
    parsed_docx = parse_uploaded_file("北京市房屋租赁合同.docx", docx_data)
    print(f"   [OK] 解析文件名: {parsed_docx['filename']}")
    print(f"   [OK] 提取标题: {parsed_docx['title']}")
    print(f"   [OK] 字符数: {parsed_docx['char_count']}")
    print(f"   [OK] 自动识别类型: {parsed_docx['contract_type']}")
    assert "租赁" in parsed_docx['contract_type'], "必须识别为租赁合同"
    
    # 测试租赁判例动态匹配
    lease_precedents = search_precedents(parsed_docx['text'])
    print(f"   [OK] 动态匹配判例: {[p['title'] for p in lease_precedents]}")
    assert any("租赁" in p['category'] or "免责" in p['category'] for p in lease_precedents)

    print("\n================================================================")
    print(">> 测试 2: 构造并测试 PDF (.pdf) 文件内存生成与解析")
    print("================================================================")
    pdf_doc = pymupdf.open()
    page = pdf_doc.new_page()
    content = "企业员工劳动聘用合同书\n\n用人单位：某某信息科技有限公司\n劳动者：赵小明\n\n第一条 合同期限为一年。其中试用期约定为六个月。\n第二条 乙方在职期间不得要求公司缴纳社会保险，如离职必须支付用人单位违约金十万元。"
    page.insert_text((50, 72), content, fontname="china-s")
    pdf_bytes = pdf_doc.write()
    pdf_doc.close()
    
    parsed_pdf = parse_uploaded_file("员工劳动合同.pdf", pdf_bytes)
    print(f"   [OK] 解析文件名: {parsed_pdf['filename']}")
    print(f"   [OK] 提取标题: {parsed_pdf['title']}")
    print(f"   [OK] 字符数: {parsed_pdf['char_count']}")
    print(f"   [OK] 自动识别类型: {parsed_pdf['contract_type']}")
    assert "劳动" in parsed_pdf['contract_type'], "必须识别为劳动合同"
    
    labor_precedents = search_precedents(parsed_pdf['text'])
    print(f"   [OK] 动态匹配判例: {[p['title'] for p in labor_precedents]}")
    assert any("劳动" in p['category'] or "试用期" in p['title'] for p in labor_precedents)

    print("\n================================================================")
    print(">> 测试 3: 测试 HTTP POST /api/contract/upload 文件上传接口")
    print("================================================================")
    # 模拟真实 HTTP 表单文件上传
    files = {
        'file': ('北京市房屋租赁合同.docx', docx_data, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    }
    r = client.post('/api/contract/upload', files=files)
    assert r.status_code == 200, f"上传接口应返回 200，实际: {r.status_code}, {r.text}"
    upload_res = r.json()['data']
    print(f"   [OK] 上传返回标题: {upload_res['title']}")
    print(f"   [OK] 上传识别合同类型: {upload_res['contract_type']}")
    print(f"   [OK] 解析内容字数: {upload_res['char_count']}")

    print("\n================================================================")
    print(">> 测试 4: 验证通用法务流式审查元提示词自适应注入")
    print("================================================================")
    from contract.prompts import UNIVERSAL_CONTRACT_AUDIT_PROMPT
    rendered_prompt = UNIVERSAL_CONTRACT_AUDIT_PROMPT.format(contract_type=upload_res['contract_type'])
    assert "房屋/物业租赁合同" in rendered_prompt
    assert "租期是否超20年法定上限" in rendered_prompt
    print("   [OK] 元提示词自适应根据租赁合同激活了民法典第705条租期上限与押金审查规则！")

    print("\n🎉 全部现场通用演示能力（Docx解析、PDF解析、HTTP上传、类型感知、法条动态命中）测试 100% 成功！")

if __name__ == "__main__":
    run_tests()
