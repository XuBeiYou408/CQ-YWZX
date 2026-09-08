"""
智审 (Doc-Agent) - 多格式合同文件离线解析与类型自动感知引擎
支持 .docx, .pdf, .txt, .md, .text 等
"""
import io
import os
from typing import Dict, Any
import docx
import pymupdf  # fitz

def extract_text_from_docx(file_bytes: bytes) -> str:
    """提取 Word (.docx) 文档全部段落与表格文本"""
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = []
    
    # 1. 提取正文段落
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            paragraphs.append(t)
            
    # 2. 提取表格文本 (合同中常有付款进度表、货物清单、交付物明细)
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                table_rows.append(" | ".join(row_text))
        if table_rows:
            paragraphs.append("【合同表格条款/清单附件】：\n" + "\n".join(table_rows))
            
    return "\n\n".join(paragraphs)

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """提取 PDF 文档全部页面文本 (使用高速本地 PyMuPDF)"""
    doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    pages_text = []
    for page in doc:
        text = page.get_text()
        if text.strip():
            pages_text.append(text.strip())
    doc.close()
    return "\n\n".join(pages_text)

def extract_text_from_plaintext(file_bytes: bytes) -> str:
    """通用纯文本解码，多编码自动回退 (UTF-8, GBK, GB18030 等)"""
    for enc in ["utf-8", "gb18030", "gbk", "utf-16", "ansi"]:
        try:
            return file_bytes.decode(enc)
        except Exception:
            continue
    return file_bytes.decode("utf-8", errors="ignore")

def detect_contract_type(text: str) -> str:
    """根据合同正文智能感知合同类型"""
    text_sample = text[:4000].lower()
    
    rules = [
        ("劳动人事用工与竞业限制协议", ["劳动合同", "聘用合同", "劳务协议", "用人单位", "劳动者", "试用期", "社会保险", "竞业限制", "辞职", "经济补偿金"]),
        ("房屋租赁与商业物业租赁合同", ["租赁合同", "出租人", "承租人", "房屋租赁", "租金", "物业费", "免租期", "押金", "转租", "铺位"]),
        ("软件技术开发与外包采购合同", ["技术开发", "定制开发", "软件开发", "源代码", "知识产权", "交付物", "验收标准", "外包", "bug", "系统上线"]),
        ("商品货物买卖与供应链采购合同", ["买卖合同", "供货合同", "采购合同", "买方", "卖方", "交付货物", "规格型号", "标的物", "质保期", "发货"]),
        ("商业保密与反不正当竞争协议 (NDA)", ["保密协议", "保密信息", "商业秘密", "披露方", "接收方", "保密义务", "知识产权", "泄密"]),
        ("民间借贷与保证担保借款合同", ["借款合同", "借款人", "出借人", "保证人", "利息", "质押", "抵押", "连带责任保证", "还款日"]),
        ("股权转让与投资合伙协议", ["股权转让", "增资扩股", "投资人", "公司章程", "持股比例", "对赌协议", "股东会", "认缴出资"]),
        ("委托服务与居间中介合同", ["委托合同", "受托人", "委托人", "咨询服务", "中介", "居间", "佣金", "服务费"]),
    ]
    
    best_type = "通用商业民商事经济合同"
    max_score = 0
    
    for c_type, keywords in rules:
        score = sum(1 for kw in keywords if kw in text_sample)
        if score > max_score and score >= 2:
            max_score = score
            best_type = c_type
            
    return best_type

def parse_uploaded_file(filename: str, content_bytes: bytes) -> Dict[str, Any]:
    """统一文件解析入口，返回纯文本与结构元数据"""
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == ".docx":
        text = extract_text_from_docx(content_bytes)
    elif ext == ".pdf":
        text = extract_text_from_pdf(content_bytes)
    else:
        text = extract_text_from_plaintext(content_bytes)
        
    text = text.strip()
    char_count = len(text)
    contract_type = detect_contract_type(text)
    
    # 提取合同标题 (前3行中长度合理的第一行)
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    title = filename
    for l in lines[:3]:
        if 4 <= len(l) <= 60 and not l.startswith("甲方") and not l.startswith("乙方"):
            title = l
            break
            
    return {
        "filename": filename,
        "title": title,
        "text": text,
        "char_count": char_count,
        "contract_type": contract_type
    }
